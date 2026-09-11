import base64
import io
import re
import warnings
import zipfile
from pathlib import Path

from docx import Document as WordDocument
from PIL import Image
from pypdf import PdfReader

from .domain import APIError

MIME = {
    "PDF": "application/pdf",
    "PNG": "image/png",
    "JPG": "image/jpeg",
    "JPEG": "image/jpeg",
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def safe_path(cfg, name):
    root = (cfg.data_dir / "files").resolve()
    path = (root / name).resolve()
    if path.parent != root:
        raise APIError("INVALID_PATH", "Invalid local file path.", 400)
    return path


def extract(path: Path, kind: str, cfg):
    """Validate content and extract actual text. Never treat filenames as trusted paths."""
    try:
        if kind == "PDF":
            if not path.read_bytes().startswith(b"%PDF-"):
                raise ValueError("signature")
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("encrypted PDF")
            if len(reader.pages) > cfg.max_document_pages:
                raise APIError(
                    "DOCUMENT_LIMIT",
                    f"Documents may contain up to {cfg.max_document_pages} pages.",
                    422,
                )
            pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        elif kind == "DOCX":
            with zipfile.ZipFile(path) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(x.file_size for x in archive.infolist()) > 64 * 1024 * 1024
                ):
                    raise ValueError("expanded document too large")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("not DOCX")
            doc = WordDocument(path)
            paragraphs = [p.text for p in doc.paragraphs]
            paragraphs.extend(
                " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            )
            pages = [(1, "\n".join(paragraphs))]
        else:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(path) as image:
                    if image.width * image.height > 25000000 or image.format not in {"JPEG", "PNG"}:
                        raise ValueError("image size or format")
                    image.verify()
            pages = [(1, "")]
        if not pages:
            raise ValueError("empty document")
        if sum(len(t) for _, t in pages) > cfg.max_document_chars:
            raise APIError(
                "DOCUMENT_LIMIT", "Document text exceeds the configured extraction limit.", 422
            )
        return pages
    except APIError:
        raise
    except Exception as exc:
        raise APIError(
            "INVALID_FILE",
            "This file is unreadable, encrypted, oversized, or does not match its extension.",
            422,
        ) from exc


def vision_images(path, kind, pages, cfg):
    wanted = [page for page, content in pages if len(content.strip()) < 20]
    if len(wanted) > cfg.max_ocr_pages:
        raise APIError(
            "OCR_PAGE_LIMIT",
            f"This file needs OCR for more than {cfg.max_ocr_pages} pages. Split it or adjust the configured limit.",
            422,
        )
    result = []
    if kind == "DOCX":
        if wanted:
            raise APIError(
                "NO_DOCUMENT_TEXT",
                "This Word document contains no extractable text. Export scanned pages as PDF or images.",
                422,
            )
        return []
    pdf = None
    try:
        if kind == "PDF":
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(path)
        for page in wanted:
            if pdf:
                pdf_page = pdf[page - 1]
                width, height = pdf_page.get_size()
                bitmap = pdf_page.render(scale=min(2, 1800 / max(width, height)))
                image = bitmap.to_pil().copy()
                bitmap.close()
                pdf_page.close()
            else:
                image = Image.open(path).convert("RGB")
            image.thumbnail((1800, 1800))
            out = io.BytesIO()
            image.convert("RGB").save(out, format="JPEG", quality=85)
            image.close()
            result.append(
                (page, "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode())
            )
    finally:
        if pdf:
            pdf.close()
    return result


def split_chunks(pages, size=1400, overlap=160):
    result = []
    for page, content in pages:
        content = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", content).strip()
        for start in range(0, len(content), size - overlap):
            text = content[start : start + size].strip()
            if text:
                result.append((page, text))
    return result
