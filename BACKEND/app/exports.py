import io
import re
from html import escape

from docx import Document
from fastapi import APIRouter, Depends, Request, Response
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .domain import APIError, audit
from .routes_tasks import authorized_job
from .security import require

router = APIRouter(tags=["Generated output"])
MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "txt": "text/plain; charset=utf-8",
}


def make_export(data, format):
    title, reply = data["title"], data.get("reply", "")
    paragraphs = ["DRAFT — HUMAN REVIEW REQUIRED", reply, "SOURCE REFERENCES"]
    paragraphs.extend(f"{c['name']} · page {c['page']}" for c in data.get("citations", []))
    out = io.BytesIO()
    if format == "docx":
        doc = Document()
        doc.add_heading(title, 0)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        doc.save(out)
    elif format == "pdf":
        style = getSampleStyleSheet()
        content = [Paragraph(escape(title), style["Title"]), Spacer(1, 15)]
        for paragraph in paragraphs:
            for line in paragraph.splitlines():
                content.append(Paragraph(escape(line) or " ", style["BodyText"]))
            content.append(Spacer(1, 10))
        SimpleDocTemplate(out, title=title, author="OMNITRIX").build(content)
    elif format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Response"
        for line in [title, *paragraphs]:
            for part in line.splitlines():
                cell = sheet.cell(sheet.max_row + 1, 1, part[:32767])
                cell.data_type = "s"  # Model or source text must never become an Excel formula.
        sheet.column_dimensions["A"].width = 100
        book.save(out)
    elif format == "pptx":
        deck = Presentation()
        for i, paragraph in enumerate(paragraphs):
            for offset in range(0, max(1, len(paragraph)), 1100):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = (
                    title if i == 0 else "Response" if i == 1 else "Source references"
                )
                slide.placeholders[1].text = paragraph[offset : offset + 1100]
        deck.save(out)
    elif format == "txt":
        out.write((title + "\n\n" + "\n\n".join(paragraphs)).encode())
    else:
        raise APIError("FORMAT", "Supported formats: docx, pdf, xlsx, pptx and txt.", 422)
    return out.getvalue()


@router.get("/tasks/{id}/outputs/{format}")
def output(id: str, format: str, request: Request, user=Depends(require("tasks"))):
    with request.app.state.db.read() as s:
        job = authorized_job(s, user, id)
        if job.status != "completed":
            raise APIError("NOT_READY", "This response is not complete yet.", 409)
        data = dict(job.data)
    result = make_export(data, format)
    filename = re.sub(r"[^a-zA-Z0-9_-]", "_", data["title"])[:80] or "omnitrix-response"
    with request.app.state.db.write() as s:
        audit(s, user, "DOCUMENT_DOWNLOADED", f"{id}/{format}", task_id=id)
    return Response(
        result,
        media_type=MIME[format],
        headers={"Content-Disposition": f'attachment; filename="{filename}.{format}"'},
    )
